import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.modules.characters.models import Character
from app.modules.characters.prompting import build_character_prompts
from app.modules.characters.schemas import CharacterCreate


@dataclass
class CharacterImportResult:
    imported_count: int = 0
    updated_count: int = 0
    skipped_count: int = 0


def import_yingge_bible(session: Session, file_path: str | Path) -> CharacterImportResult:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("导入角色圣经需要安装 openpyxl。") from exc

    path = Path(file_path)
    workbook = openpyxl.load_workbook(path, data_only=True)
    sheet = workbook["角色IP圣经"] if "角色IP圣经" in workbook.sheetnames else workbook.worksheets[0]
    result = CharacterImportResult()
    for row_index in range(4, sheet.max_row + 1):
        raw_identity = _clean(sheet.cell(row_index, 1).value)
        if not raw_identity:
            result.skipped_count += 1
            continue
        payload = _row_to_payload(sheet, row_index, raw_identity)
        existing = session.scalar(select(Character).where(Character.name == payload.name, Character.source == "yingge_bible"))
        data = payload.model_dump()
        data.update({key: value for key, value in build_character_prompts(payload).items() if not data.get(key)})
        if existing:
            for key, value in data.items():
                setattr(existing, key, value)
            result.updated_count += 1
        else:
            session.add(Character(**data))
            result.imported_count += 1
    session.commit()
    return result


def _row_to_payload(sheet: Any, row_index: int, raw_identity: str) -> CharacterCreate:
    return CharacterCreate(
        name=_extract(raw_identity, "姓名") or raw_identity.splitlines()[0].replace("姓名：", "").strip(),
        alias=_extract(raw_identity, "绰号"),
        rank=_extract(raw_identity, "排名"),
        star=_extract(raw_identity, "星位"),
        origin=_extract(raw_identity, "出身"),
        ip_name="英歌水浒",
        role_type=_infer_role_type(_clean(sheet.cell(row_index, 9).value)),
        bio=raw_identity,
        liangshan_role=_clean(sheet.cell(row_index, 2).value),
        weapons=_clean(sheet.cell(row_index, 3).value),
        weapon=_clean(sheet.cell(row_index, 3).value),
        personality_tags=_clean(sheet.cell(row_index, 4).value),
        internal_conflict=_clean(sheet.cell(row_index, 5).value),
        life_events=_clean(sheet.cell(row_index, 6).value),
        audience_hook=_clean(sheet.cell(row_index, 7).value),
        target_audience=_clean(sheet.cell(row_index, 8).value),
        yingge_role=_clean(sheet.cell(row_index, 9).value),
        facepaint_main_color=_clean(sheet.cell(row_index, 10).value),
        color_symbolism=_clean(sheet.cell(row_index, 11).value),
        facepaint_patterns=_clean(sheet.cell(row_index, 12).value),
        source_color_clues=_clean(sheet.cell(row_index, 13).value),
        visual_tone_keywords=_clean(sheet.cell(row_index, 14).value),
        positive_prompt_terms=_clean(sheet.cell(row_index, 15).value),
        negative_prompt_terms=_clean(sheet.cell(row_index, 16).value),
        visual_keywords=_clean(sheet.cell(row_index, 15).value),
        negative_keywords=_clean(sheet.cell(row_index, 16).value),
        narrative_origin=_clean(sheet.cell(row_index, 17).value),
        relationship_map=_clean(sheet.cell(row_index, 18).value),
        product_tone=_clean(sheet.cell(row_index, 19).value),
        blessing_meaning=_clean(sheet.cell(row_index, 20).value),
        source="yingge_bible",
        source_row=row_index,
        status="active",
    )


def _extract(text: str, label: str) -> str | None:
    match = re.search(rf"{label}[:：]([^\n]+)", text)
    return match.group(1).strip() if match else None


def _clean(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _infer_role_type(yingge_role: str | None) -> str | None:
    if not yingge_role:
        return None
    for keyword in ["司鼓者", "指挥", "旗手", "引舞者", "八头槌", "槌手", "女将"]:
        if keyword in yingge_role:
            return keyword
    return "英歌角色"
